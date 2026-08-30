import os
from PyPDF2 import PdfReader
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook


# ==========================================
# Universal File Reader
# ==========================================

def read_file(filename):

    ext = os.path.splitext(filename)[1].lower()

    if ext == ".txt":
        try:
            with open(filename, "r", encoding="utf-8", errors="ignore") as file:
                return file.read()
        except Exception as e:
            return str(e)

    elif ext == ".pdf":
        return read_pdf(filename)

    elif ext == ".docx":
        return read_docx(filename)

    elif ext == ".pptx":
        return read_pptx(filename)

    elif ext == ".xlsx":
        return read_excel(filename)

    else:
        return "Unsupported file type."


# ==========================================
# TXT Reader
# ==========================================

def read_all_txt(path="."):

    try:

        result = ""

        for file in os.listdir(path):

            if file.endswith(".txt"):

                with open(file, "r", encoding="utf-8", errors="ignore") as f:

                    result += f"\n--- {file} ---\n"
                    result += f.read() + "\n"

        return result if result else "No txt files found."

    except Exception as e:

        return str(e)


# ==========================================
# List Files
# ==========================================

def list_files(path="."):

    try:
        return "\n".join(os.listdir(path))

    except Exception as e:
        return str(e)


# ==========================================
# PDF Reader
# ==========================================

def read_pdf(filename):

    try:

        reader = PdfReader(filename)

        text = ""

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:
                text += page_text + "\n"

        return text if text.strip() else "No text found in PDF."

    except Exception as e:

        return f"PDF Error: {e}"


# ==========================================
# DOCX Reader
# ==========================================

def read_docx(filename):

    try:

        doc = Document(filename)

        text = ""

        for para in doc.paragraphs:
            text += para.text + "\n"

        return text if text.strip() else "Document is empty."

    except Exception as e:

        return f"DOCX Error: {e}"


# ==========================================
# PPTX Reader
# ==========================================

def read_pptx(filename):

    try:

        prs = Presentation(filename)

        text = ""

        for slide_num, slide in enumerate(prs.slides, start=1):

            text += f"\n--- Slide {slide_num} ---\n"

            for shape in slide.shapes:

                if hasattr(shape, "text"):
                    text += shape.text + "\n"

        return text if text.strip() else "Presentation is empty."

    except Exception as e:

        return f"PPTX Error: {e}"


# ==========================================
# Excel Reader
# ==========================================

def read_excel(filename):

    try:

        wb = load_workbook(filename, data_only=True)

        result = ""

        for sheet in wb.sheetnames:

            result += f"\n--- Sheet: {sheet} ---\n"

            ws = wb[sheet]

            for row in ws.iter_rows(values_only=True):

                line = "\t".join(str(cell) for cell in row)

                result += line + "\n"

        return result

    except Exception as e:

        return f"Excel Error: {e}"